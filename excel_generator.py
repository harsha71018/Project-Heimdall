import os
import re
import json
import glob
import tempfile
from datetime import datetime
from typing import Dict, Any, Optional, List, Tuple
from urllib.parse import urlsplit, urlunsplit, parse_qsl, urlencode

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
import pandas as pd

# Query parameters that uniquely identify a job posting on aggregator/ATS platforms
IDENTITY_QUERY_PARAMS = {
    "jk", "jl", "id", "jobreferencecode", "jobid", "page", "job"
}


def canonicalize_job_url(url: str) -> str:
    """
    Normalizes a job URL for ledger and deduplication lookups while preserving
    platform-required identity query parameters (e.g., Indeed ?jk=, Glassdoor ?jl=, ATS IDs).
    """
    if not url or not isinstance(url, str):
        return ""
    raw = url.strip()
    if not raw.lower().startswith(("http://", "https://")):
        return ""
    try:
        parsed = urlsplit(raw)
        if not parsed.scheme or not parsed.hostname:
            return ""
        kept_params = [
            (k.lower(), v.strip().lower())
            for k, v in parse_qsl(parsed.query, keep_blank_values=False)
            if k.lower() in IDENTITY_QUERY_PARAMS
        ]
        kept_params.sort()
        clean_path = parsed.path.rstrip("/").lower()
        return urlunsplit((
            parsed.scheme.lower(),
            parsed.hostname.lower(),
            clean_path,
            urlencode(kept_params),
            ""
        ))
    except Exception:
        return re.sub(r'\?.*$', '', raw).strip().lower()


class ExcelDashboardGenerator:
    """
    Builds a styled, interactive Excel Workbook for job application tracking.
    Features:
    - Native openpyxl hyperlinks (handles long UTM tracking URLs safely)
    - Persistent Application Status & Notes memory via applications.json (atomic writes)
    - Archive-safe sync: leads/ sorted by filename, Tracker_Latest*.xlsx always takes final precedence
    - Dedicated 'Applied & In Progress' tab preserving active applications across days
    - Tab 'Run Summary' audit dashboard
    - Formula injection sanitization for untrusted strings
    """

    APPLIED_TAB_NAME = "Applied & In Progress"

    def __init__(self, config: Dict[str, Any]):
        self.output_config = config.get("output", {})
        self.filename = self.output_config.get("filename", "Job_Applications_Tracker.xlsx")
        self.sheet_name = self.output_config.get("sheet_name", "Active Job Leads")
        self.status_options = self.output_config.get("status_options", [
            "Not Applied", "Applied", "Referral Requested",
            "Screening Call", "Interviewing", "Assessment", "Rejected", "Offer Received"
        ])
        self.leads_dir = self.output_config.get("leads_dir", "leads")
        self.stable_filename = self.output_config.get("stable_filename", "Tracker_Latest.xlsx")
        self.applications_file = self.output_config.get("applications_file", "applications.json")
        self.applications_ledger = self._sync_and_load_applications()

    @staticmethod
    def _sanitize_cell_value(val: Any) -> str:
        """Neutralizes spreadsheet formula injection on untrusted strings."""
        if val is None:
            return ""
        s = str(val)
        if s and s != "-" and s[0] in ("=", "+", "-", "@"):
            return f"'{s}"
        return s

    @staticmethod
    def _unescape_cell_value(val: Any) -> str:
        """Strips leading formula-escape single quote when syncing back from Excel."""
        if val is None:
            return ""
        s = str(val).strip()
        if len(s) > 1 and s[0] == "'" and s[1] in ("=", "+", "-", "@"):
            return s[1:].strip()
        return s

    def _atomic_save_applications(self, ledger: Dict[str, Dict[str, str]]) -> None:
        """Atomically writes applications.json using a temp file and os.replace."""
        abs_target = os.path.abspath(self.applications_file)
        parent_dir = os.path.dirname(abs_target)
        os.makedirs(parent_dir, exist_ok=True)
        tmp_path = None
        try:
            with tempfile.NamedTemporaryFile(
                mode="w", dir=parent_dir, delete=False, encoding="utf-8", suffix=".tmp"
            ) as tmp_f:
                tmp_path = tmp_f.name
                json.dump(ledger, tmp_f, indent=2)
                tmp_f.flush()
                os.fsync(tmp_f.fileno())
            os.replace(tmp_path, abs_target)
        except Exception as e:
            print(f"[!] Warning: Could not save '{self.applications_file}': {e}")
        finally:
            if tmp_path and os.path.exists(tmp_path):
                try:
                    os.remove(tmp_path)
                except OSError:
                    pass

    def _get_ordered_sync_candidates(self) -> List[Tuple[str, bool]]:
        """
        Gathers candidate tracker workbooks as (filepath, is_primary_tracker) tuples:
        1. Historical archives in leads/Job_Leads_*.xlsx sorted strictly by filename
           (so opening an old archive in Excel and touching its mtime never overrides newer runs).
        2. Primary trackers (Tracker_Latest*.xlsx and filename) sorted by mtime and placed LAST
           so edits in Tracker_Latest*.xlsx always win and can revert statuses to 'Not Applied'.
        """
        archive_candidates = []
        primary_candidates = []
        seen_abs = set()

        if os.path.isdir(self.leads_dir):
            pattern = os.path.join(self.leads_dir, "Job_Leads_*.xlsx")
            lead_files = sorted(glob.glob(pattern), key=lambda p: os.path.basename(p))
            for p in lead_files:
                abs_p = os.path.abspath(p)
                if os.path.exists(abs_p) and abs_p not in seen_abs:
                    seen_abs.add(abs_p)
                    archive_candidates.append((p, False))

        primary_raw = []
        if self.stable_filename:
            base, ext = os.path.splitext(self.stable_filename)
            for p in glob.glob(f"{base}*{ext}"):
                if os.path.exists(p):
                    primary_raw.append(p)
            if os.path.exists(self.stable_filename) and self.stable_filename not in primary_raw:
                primary_raw.append(self.stable_filename)

        if self.filename and os.path.exists(self.filename):
            primary_raw.append(self.filename)

        def _mtime_key(path: str):
            try:
                return (round(os.path.getmtime(path), 3), os.path.basename(path))
            except OSError:
                return (0.0, os.path.basename(path))

        primary_raw.sort(key=_mtime_key)
        for p in primary_raw:
            abs_p = os.path.abspath(p)
            if abs_p not in seen_abs:
                seen_abs.add(abs_p)
                primary_candidates.append((p, True))

        return archive_candidates + primary_candidates

    def _sync_sheet_into_ledger(
        self,
        ws,
        ledger: Dict[str, Dict[str, str]],
        allow_revert: bool = False,
        baseline: Optional[Dict[str, Dict[str, str]]] = None,
        initial_ledger_keys: Optional[set] = None
    ) -> int:
        """Extracts status, notes, and job metadata from a worksheet into ledger."""
        if baseline is None:
            baseline = {}
        if initial_ledger_keys is None:
            initial_ledger_keys = set()

        headers = [str(cell.value or "").strip() for cell in ws[1]]
        header_indices = {}
        for idx, h in enumerate(headers, 1):
            h_lower = h.lower()
            if "application status" in h_lower:
                header_indices["status"] = idx
            elif "referral / notes" in h_lower or "notes" in h_lower:
                header_indices["notes"] = idx
            elif "apply action" in h_lower or "apply" in h_lower:
                header_indices["apply"] = idx
            elif "company" in h_lower:
                header_indices["company"] = idx
            elif "title" in h_lower:
                header_indices["title"] = idx
            elif "location" in h_lower:
                header_indices["location"] = idx
            elif "platform" in h_lower:
                header_indices["platform"] = idx

        status_col = header_indices.get("status")
        notes_col = header_indices.get("notes")
        apply_col = header_indices.get("apply")
        company_col = header_indices.get("company")
        title_col = header_indices.get("title")
        location_col = header_indices.get("location")
        platform_col = header_indices.get("platform")

        if not (status_col or notes_col):
            return 0

        updates = 0
        for row in range(2, ws.max_row + 1):
            raw_status = ws.cell(row=row, column=status_col).value if status_col else ""
            raw_notes = ws.cell(row=row, column=notes_col).value if notes_col else ""

            status_val = self._unescape_cell_value(raw_status)
            notes_val = self._unescape_cell_value(raw_notes)

            apply_cell = ws.cell(row=row, column=apply_col) if apply_col else None
            url_target = ""
            if apply_cell:
                if apply_cell.hyperlink and apply_cell.hyperlink.target:
                    url_target = str(apply_cell.hyperlink.target).strip()
                elif apply_cell.value and str(apply_cell.value).strip().startswith("http"):
                    url_target = str(apply_cell.value).strip()

            norm_url = canonicalize_job_url(url_target) if url_target else ""
            company_orig = self._unescape_cell_value(
                ws.cell(row=row, column=company_col).value if company_col else ""
            )
            title_orig = self._unescape_cell_value(
                ws.cell(row=row, column=title_col).value if title_col else ""
            )
            location_orig = self._unescape_cell_value(
                ws.cell(row=row, column=location_col).value if location_col else ""
            )
            platform_orig = self._unescape_cell_value(
                ws.cell(row=row, column=platform_col).value if platform_col else ""
            )

            company_val = company_orig.lower()
            title_val = title_orig.lower()

            if norm_url:
                key = norm_url
            elif company_val and title_val and company_val not in ("no jobs found", "no active applications yet"):
                key = f"{company_val}::{title_val}"
            else:
                key = ""

            if not key:
                continue

            # Fix N2: Old archives in leads/ cannot resurrect or modify keys already in applications.json
            if not allow_revert and key in initial_ledger_keys:
                continue

            eff_status = status_val if (status_val and status_val != "None") else "Not Applied"
            eff_notes = notes_val or ""

            # Fix N1: If cell matches pre-sync baseline, skip so unchanged tabs don't stomp active edits
            if key in baseline:
                base_status = baseline[key].get("status", "Not Applied")
                base_notes = baseline[key].get("notes", "")
                if eff_status == base_status and eff_notes == base_notes:
                    continue

            is_default_empty = (status_val in ["", "Not Applied", "None"]) and not notes_val
            if is_default_empty:
                if allow_revert and key in ledger:
                    prev_status = ledger[key].get("status", "Not Applied")
                    prev_notes = ledger[key].get("notes", "")
                    if prev_status != "Not Applied" or prev_notes != "":
                        existing_meta = ledger[key]
                        ledger[key] = {
                            "status": "Not Applied",
                            "notes": "",
                            "company": company_orig or existing_meta.get("company", ""),
                            "title": title_orig or existing_meta.get("title", ""),
                            "location": location_orig or existing_meta.get("location", ""),
                            "platform": platform_orig or existing_meta.get("platform", ""),
                            "url": url_target or existing_meta.get("url", ""),
                            "last_updated": datetime.now().strftime("%Y-%m-%d")
                        }
                        updates += 1
                continue

            clean_status = status_val if status_val and status_val != "None" else "Not Applied"
            existing_meta = ledger.get(key, {})
            ledger[key] = {
                "status": clean_status,
                "notes": notes_val,
                "company": company_orig or existing_meta.get("company", ""),
                "title": title_orig or existing_meta.get("title", ""),
                "location": location_orig or existing_meta.get("location", ""),
                "platform": platform_orig or existing_meta.get("platform", ""),
                "url": url_target or existing_meta.get("url", ""),
                "last_updated": datetime.now().strftime("%Y-%m-%d")
            }
            updates += 1

        return updates

    def _sync_and_load_applications(self) -> Dict[str, Dict[str, str]]:
        """
        Loads existing applications.json memory as baseline source of truth, then
        syncs newly recorded statuses/notes across historical archives (sorted by filename)
        followed by Tracker_Latest*.xlsx (which always has final precedence and revert rights).
        """
        ledger: Dict[str, Dict[str, str]] = {}
        if os.path.exists(self.applications_file):
            try:
                with open(self.applications_file, "r", encoding="utf-8") as f:
                    data = json.load(f)
                    if isinstance(data, dict):
                        ledger = data
            except Exception as e:
                print(f"[!] Warning: Could not read '{self.applications_file}': {e}")

        initial_ledger_keys = set(ledger.keys())
        candidate_files = self._get_ordered_sync_candidates()
        total_updates = 0

        for target_file, is_primary in candidate_files:
            wb = None
            try:
                wb = openpyxl.load_workbook(target_file, data_only=True)
                sheets_to_read = []
                if self.sheet_name in wb.sheetnames:
                    sheets_to_read.append(wb[self.sheet_name])
                else:
                    sheets_to_read.append(wb.active)

                if self.APPLIED_TAB_NAME in wb.sheetnames:
                    sheets_to_read.append(wb[self.APPLIED_TAB_NAME])

                baseline = {k: dict(v) for k, v in ledger.items()}

                for ws in sheets_to_read:
                    total_updates += self._sync_sheet_into_ledger(
                        ws,
                        ledger,
                        allow_revert=is_primary,
                        baseline=baseline,
                        initial_ledger_keys=initial_ledger_keys
                    )

            except PermissionError:
                print(f"[!] Warning: Tracker file '{target_file}' is currently locked in Excel.")
                print(f"    Proceeding safely with existing '{self.applications_file}' memory.")
            except Exception as e:
                print(f"[!] Warning: Could not sync from '{target_file}': {e}")
            finally:
                if wb is not None:
                    try:
                        wb.close()
                    except Exception:
                        pass

        if total_updates > 0:
            self._atomic_save_applications(ledger)

        return ledger

    def generate(self, jobs_df: pd.DataFrame, output_path: str = None, run_stats: dict = None) -> str:
        """
        Creates and styles the Excel spreadsheet from the enriched DataFrame.
        Produces:
          - Tab 1: Active Job Leads (today's scraped leads with status/notes carried over)
          - Tab 2: Applied & In Progress (permanent tracker of all active applications in applications.json)
          - Tab 3: Run Summary (detailed funnel audit statistics)
          - Dual outputs: timestamped archive in leads/ and stable Tracker_Latest.xlsx in root.
        """
        save_stable = False
        if output_path is not None:
            primary_target = output_path
        else:
            os.makedirs(self.leads_dir, exist_ok=True)
            now_str = datetime.now().strftime("%Y-%m-%d_%H%M")
            primary_target = os.path.join(self.leads_dir, f"Job_Leads_{now_str}.xlsx")
            save_stable = True

        wb = openpyxl.Workbook()
        try:
            ws = wb.active
            ws.title = self.sheet_name
            ws.views.sheetView[0].showGridLines = True

            headers = [
                "Sl. No",
                "Lead Status",
                "Date Posted",
                "Company Name",
                "Job Title / Role",
                "Main Skills Required",
                "Experience Level",
                "Location / Mode",
                "Platform",
                "Apply Action",
                "Application Status",
                "Referral / Notes"
            ]
            ws.append(headers)

            # Pre-allocated Theme Colors, Fonts & Alignments
            header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
            header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")

            row_alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
            row_white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

            apply_btn_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
            apply_btn_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF", underline="single")

            skill_font = Font(name="Segoe UI Semibold", size=10, color="047857")
            standard_font = Font(name="Segoe UI", size=10, color="1E293B")
            company_font = Font(name="Segoe UI", size=10, bold=True, color="0F172A")
            applied_font = Font(name="Segoe UI Semibold", size=10, color="2563EB")
            new_lead_font = Font(name="Segoe UI", size=10, bold=True, color="047857")
            seen_lead_font = Font(name="Segoe UI", size=10, color="64748B")
            summary_val_font = Font(name="Segoe UI Semibold", size=10, color="0F172A")

            align_center = Alignment(horizontal="center", vertical="center")
            align_header = Alignment(horizontal="center", vertical="center", wrap_text=True)
            align_left = Alignment(vertical="center")
            align_indent = Alignment(vertical="center", indent=1)

            thin_border = Border(
                left=Side(style='thin', color='E2E8F0'),
                right=Side(style='thin', color='E2E8F0'),
                top=Side(style='thin', color='E2E8F0'),
                bottom=Side(style='thin', color='E2E8F0')
            )

            # Style Header Row for Tab 1
            ws.row_dimensions[1].height = 28
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=col_idx)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = align_header
                cell.border = thin_border

            ledger_metadata_updated = False

            if jobs_df is None or jobs_df.empty:
                ws.append([
                    1, "N/A", "N/A", "No Jobs Found",
                    "Try adjusting search query or location",
                    "-", "-", "-", "-", "No Link", "Not Applied", ""
                ])
            else:
                for row_num, (_, row) in enumerate(jobs_df.iterrows(), 1):
                    row_idx = ws.max_row + 1
                    ws.row_dimensions[row_idx].height = 24

                    is_even = (row_idx % 2 == 0)
                    current_fill = row_alt_fill if is_even else row_white_fill

                    lead_status = str(row.get("lead_status", "🟢 NEW"))
                    date_posted = str(row.get("formatted_date", "Recent"))
                    company_raw = str(row.get("company", "Hiring Company"))
                    title_raw = str(row.get("title", "Job Role"))
                    skills_raw = str(row.get("extracted_skills", "General Technical Skills"))
                    exp_raw = str(row.get("extracted_experience", "Not Specified"))
                    work_mode_raw = str(row.get("work_mode", row.get("location", "Hyderabad, India")))
                    platform_raw = str(row.get("site", "Web")).capitalize()
                    apply_url = str(row.get("final_apply_url", "")).strip()

                    # Memory lookup: check canonical URL or company::title in applications_ledger
                    norm_url = canonicalize_job_url(apply_url) if apply_url else ""
                    fallback_key = f"{company_raw.strip().lower()}::{title_raw.strip().lower()}"
                    matched_key = (
                        norm_url if (norm_url and norm_url in self.applications_ledger)
                        else (fallback_key if fallback_key in self.applications_ledger else None)
                    )
                    saved_entry = self.applications_ledger.get(matched_key) if matched_key else None

                    if saved_entry:
                        status_val = saved_entry.get("status", "Not Applied")
                        notes_val = saved_entry.get("notes", "")
                        # Backfill metadata in ledger if missing
                        if not saved_entry.get("company") or not saved_entry.get("url"):
                            saved_entry["company"] = company_raw
                            saved_entry["title"] = title_raw
                            saved_entry["location"] = work_mode_raw
                            saved_entry["platform"] = platform_raw
                            saved_entry["url"] = apply_url
                            ledger_metadata_updated = True
                    else:
                        status_val = "Not Applied"
                        notes_val = ""

                    c1 = ws.cell(row=row_idx, column=1, value=row_num)
                    c2 = ws.cell(row=row_idx, column=2, value=self._sanitize_cell_value(lead_status))
                    c3 = ws.cell(row=row_idx, column=3, value=self._sanitize_cell_value(date_posted))
                    c4 = ws.cell(row=row_idx, column=4, value=self._sanitize_cell_value(company_raw))
                    c5 = ws.cell(row=row_idx, column=5, value=self._sanitize_cell_value(title_raw))
                    c6 = ws.cell(row=row_idx, column=6, value=self._sanitize_cell_value(skills_raw))
                    c7 = ws.cell(row=row_idx, column=7, value=self._sanitize_cell_value(exp_raw))
                    c8 = ws.cell(row=row_idx, column=8, value=self._sanitize_cell_value(work_mode_raw))
                    c9 = ws.cell(row=row_idx, column=9, value=self._sanitize_cell_value(platform_raw))
                    c10 = ws.cell(row=row_idx, column=10)
                    c11 = ws.cell(row=row_idx, column=11, value=self._sanitize_cell_value(status_val))
                    c12 = ws.cell(row=row_idx, column=12, value=self._sanitize_cell_value(notes_val))

                    for c in [c1, c2, c3, c4, c5, c6, c7, c8, c9, c11, c12]:
                        c.fill = current_fill
                        c.border = thin_border
                        c.alignment = align_left

                    c1.alignment = align_center
                    c2.font = new_lead_font if "NEW" in lead_status else seen_lead_font
                    c2.alignment = align_center
                    c3.alignment = align_center
                    c4.font = company_font
                    c5.font = standard_font
                    c6.font = skill_font
                    c7.alignment = align_center
                    c9.alignment = align_center
                    c11.alignment = align_center
                    c11.font = applied_font if (status_val and status_val != "Not Applied") else standard_font
                    c12.font = standard_font

                    # Native openpyxl Hyperlink
                    c10.border = thin_border
                    if apply_url and apply_url.lower().startswith(("http://", "https://")):
                        c10.value = "Apply Now"
                        c10.hyperlink = apply_url
                        c10.font = apply_btn_font
                        c10.fill = apply_btn_fill
                        c10.alignment = align_center
                    else:
                        c10.value = "Link Unavailable"
                        c10.font = standard_font
                        c10.fill = current_fill
                        c10.alignment = align_center

            if ledger_metadata_updated:
                self._atomic_save_applications(self.applications_ledger)

            # Add Status Dropdown Data Validation (Column K on Tab 1)
            status_list_str = f'"{",".join(self.status_options)}"'
            if ws.max_row >= 2:
                status_dv = DataValidation(
                    type="list",
                    formula1=status_list_str,
                    allow_blank=True,
                    showErrorMessage=True
                )
                status_dv.error = "Please select a valid application status"
                status_dv.errorTitle = "Invalid Status"
                ws.add_data_validation(status_dv)
                status_dv.add(f"K2:K{ws.max_row}")

            last_col_letter = get_column_letter(len(headers))
            ws.auto_filter.ref = f"A1:{last_col_letter}{ws.max_row}"
            ws.freeze_panes = "A2"

            col_widths = {
                1: 8, 2: 15, 3: 14, 4: 26, 5: 34, 6: 35,
                7: 22, 8: 22, 9: 15, 10: 18, 11: 20, 12: 28
            }
            for col_idx, width in col_widths.items():
                ws.column_dimensions[get_column_letter(col_idx)].width = width

            # -----------------------------------------------------------------
            # Build Tab 2: Applied & In Progress (Permanent Kanban View)
            # -----------------------------------------------------------------
            ws_applied = wb.create_sheet(title=self.APPLIED_TAB_NAME)
            ws_applied.views.sheetView[0].showGridLines = True
            applied_headers = [
                "Sl. No",
                "Company Name",
                "Job Title / Role",
                "Location / Mode",
                "Platform",
                "Apply Action",
                "Application Status",
                "Referral / Notes",
                "Last Updated"
            ]
            ws_applied.append(applied_headers)
            ws_applied.row_dimensions[1].height = 28

            for col_idx in range(1, len(applied_headers) + 1):
                cell = ws_applied.cell(row=1, column=col_idx)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = align_header
                cell.border = thin_border

            active_apps = [
                (k, v) for k, v in self.applications_ledger.items()
                if isinstance(v, dict) and (
                    v.get("status", "Not Applied") not in ("", "Not Applied", "None")
                    or bool(str(v.get("notes", "")).strip())
                )
            ]

            if not active_apps:
                ws_applied.append([
                    1, "No Active Applications Yet",
                    "Mark jobs as 'Applied' in Tab 1 to track them permanently here",
                    "-", "-", "No Link", "Not Applied", "", "-"
                ])
            else:
                for app_idx, (key, entry) in enumerate(active_apps, 1):
                    r_idx = ws_applied.max_row + 1
                    ws_applied.row_dimensions[r_idx].height = 24
                    c_fill = row_alt_fill if (r_idx % 2 == 0) else row_white_fill

                    co_str = entry.get("company") or (key.split("::")[0].title() if "::" in key else "Tracked Company")
                    ti_str = entry.get("title") or (key.split("::")[1].title() if "::" in key else "Tracked Role")
                    loc_str = entry.get("location") or "-"
                    plat_str = entry.get("platform") or "-"
                    url_str = entry.get("url") or (key if key.startswith("http") else "")
                    st_str = entry.get("status", "Applied")
                    nt_str = entry.get("notes", "")
                    upd_str = entry.get("last_updated", "-")

                    ca1 = ws_applied.cell(row=r_idx, column=1, value=app_idx)
                    ca2 = ws_applied.cell(row=r_idx, column=2, value=self._sanitize_cell_value(co_str))
                    ca3 = ws_applied.cell(row=r_idx, column=3, value=self._sanitize_cell_value(ti_str))
                    ca4 = ws_applied.cell(row=r_idx, column=4, value=self._sanitize_cell_value(loc_str))
                    ca5 = ws_applied.cell(row=r_idx, column=5, value=self._sanitize_cell_value(plat_str))
                    ca6 = ws_applied.cell(row=r_idx, column=6)
                    ca7 = ws_applied.cell(row=r_idx, column=7, value=self._sanitize_cell_value(st_str))
                    ca8 = ws_applied.cell(row=r_idx, column=8, value=self._sanitize_cell_value(nt_str))
                    ca9 = ws_applied.cell(row=r_idx, column=9, value=self._sanitize_cell_value(upd_str))

                    for c in [ca1, ca2, ca3, ca4, ca5, ca7, ca8, ca9]:
                        c.fill = c_fill
                        c.border = thin_border
                        c.alignment = align_left

                    ca1.alignment = align_center
                    ca2.font = company_font
                    ca3.font = standard_font
                    ca5.alignment = align_center
                    ca7.alignment = align_center
                    ca7.font = applied_font if st_str != "Not Applied" else standard_font
                    ca8.font = standard_font
                    ca9.alignment = align_center
                    ca9.font = standard_font

                    ca6.border = thin_border
                    if url_str and url_str.lower().startswith(("http://", "https://")):
                        ca6.value = "View Job"
                        ca6.hyperlink = url_str
                        ca6.font = apply_btn_font
                        ca6.fill = apply_btn_fill
                        ca6.alignment = align_center
                    else:
                        ca6.value = "Link Unavailable"
                        ca6.font = standard_font
                        ca6.fill = c_fill
                        ca6.alignment = align_center

            if ws_applied.max_row >= 2:
                applied_dv = DataValidation(
                    type="list",
                    formula1=status_list_str,
                    allow_blank=True,
                    showErrorMessage=True
                )
                ws_applied.add_data_validation(applied_dv)
                applied_dv.add(f"G2:G{ws_applied.max_row}")

            ws_applied.auto_filter.ref = f"A1:{get_column_letter(len(applied_headers))}{ws_applied.max_row}"
            ws_applied.freeze_panes = "A2"
            applied_widths = {1: 8, 2: 26, 3: 34, 4: 22, 5: 15, 6: 18, 7: 20, 8: 30, 9: 16}
            for col_idx, width in applied_widths.items():
                ws_applied.column_dimensions[get_column_letter(col_idx)].width = width

            # -----------------------------------------------------------------
            # Build Tab 3: Run Summary
            # -----------------------------------------------------------------
            ws_summary = wb.create_sheet(title="Run Summary")
            ws_summary.views.sheetView[0].showGridLines = True
            ws_summary.append(["Audit Metric", "Value"])
            ws_summary.row_dimensions[1].height = 26

            for col_idx in [1, 2]:
                cell = ws_summary.cell(row=1, column=col_idx)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = align_center
                cell.border = thin_border

            stats = run_stats or {}
            mode_val = str(stats.get("mode", "frugal")).upper()
            total_leads_val = len(jobs_df) if jobs_df is not None else 0
            fc_cands = int(stats.get("fc_candidates", 0) or 0)
            if "jobspy_candidates" in stats:
                jobspy_cands = int(stats.get("jobspy_candidates", 0) or 0)
            else:
                jobspy_cands = int(stats.get("jobspy_retained", total_leads_val) or 0)
            total_raw_candidates = fc_cands + jobspy_cands

            credits_val = stats.get("credits_used", 0)
            sr_dropped = stats.get("dropped_senior_title", 0)
            ceil_dropped = stats.get("dropped_experience_ceiling", 0)
            ts_val = stats.get("timestamp", datetime.now().strftime("%Y-%m-%d %H:%M"))

            summary_rows = [
                ("Run Timestamp", ts_val),
                ("Scraping Mode", mode_val),
                ("Total Raw Candidates Gathered", total_raw_candidates),
                ("Senior Roles Filtered", sr_dropped),
                ("Experience Ceiling Filtered (> 2 YOE)", ceil_dropped),
                ("Final Verified Leads Retained", total_leads_val),
                ("Firecrawl Credits Consumed", credits_val)
            ]

            for r_idx, (metric, val) in enumerate(summary_rows, 2):
                ws_summary.row_dimensions[r_idx].height = 22
                c_m = ws_summary.cell(row=r_idx, column=1, value=metric)
                c_v = ws_summary.cell(row=r_idx, column=2, value=val)

                c_fill = row_alt_fill if (r_idx % 2 == 0) else row_white_fill
                c_m.fill = c_fill
                c_m.font = standard_font
                c_m.border = thin_border
                c_m.alignment = align_indent

                c_v.fill = c_fill
                c_v.font = summary_val_font
                c_v.border = thin_border
                c_v.alignment = align_center

            ws_summary.column_dimensions["A"].width = 40
            ws_summary.column_dimensions["B"].width = 24

            saved_file = self._safe_save(wb, primary_target)
            if save_stable:
                saved_stable = self._safe_save(wb, self.stable_filename)
                try:
                    st = os.stat(saved_file)
                    os.utime(saved_stable, (st.st_atime, st.st_mtime))
                except OSError:
                    pass

            print(f"\n[+] Excel Workbook successfully created: {os.path.abspath(saved_file)}")
            return saved_file
        finally:
            try:
                wb.close()
            except Exception:
                pass

    def _safe_save(self, wb, output_path: str) -> str:
        """Saves workbook safely, falling back if the target or fallback file is locked by Excel."""
        parent = os.path.dirname(os.path.abspath(output_path))
        if parent:
            os.makedirs(parent, exist_ok=True)
        try:
            wb.save(output_path)
            return output_path
        except PermissionError:
            base, ext = os.path.splitext(output_path)
            fallback_path = f"{base}_Latest{ext}"
            print(f"\n[!] Warning: '{output_path}' is currently open in Excel.")
            print(f"    Saving clean copy to: '{fallback_path}' instead.")
            try:
                wb.save(fallback_path)
                return fallback_path
            except PermissionError:
                ts_fallback = f"{base}_{datetime.now().strftime('%Y%m%d_%H%M%S')}{ext}"
                wb.save(ts_fallback)
                return ts_fallback
